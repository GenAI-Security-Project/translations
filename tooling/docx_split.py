"""Split a Heading-1-structured .docx into one Markdown file per section,
plus one SVG "figure" per embedded image that turns out to carry translatable
text (see image_svg.py for the OCR/blank/overlay mechanism).

Used by translate_section.py the first time an asset whose registry
`split_by` is `heading_1` is processed for a release — once per release
(operating on _source/), never once per locale.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterator, List, Tuple

from docx import Document
from docx.oxml.ns import qn
from text_quality import is_footer_url_artifact, is_garbled, is_page_number_artifact
from docx.text.paragraph import Paragraph

import image_svg
from source_manifest import write_manifest


def _slugify(heading: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", heading).strip("_")
    return slug or "Section"


def _iter_all_paragraphs(doc: Document) -> Iterator[Paragraph]:
    """doc.paragraphs only sees direct body children — it misses paragraphs
    wrapped in <w:sdt> (structured document tags), which real-world exports
    (observed: a Google-Docs-produced .docx) use freely. Walk the full tree
    instead so every heading is found regardless of what wraps it."""
    for p_element in doc.element.body.iter(qn("w:p")):
        yield Paragraph(p_element, doc)


_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"  # legacy VML images; not in python-docx's default nsmap


def _image_blobs_in(doc: Document, p_element) -> List[bytes]:
    """Any images embedded directly in this paragraph, as raw bytes."""
    rids = [blip.get(qn("r:embed")) for blip in p_element.iter(qn("a:blip"))]
    rids += [img.get(qn("r:id")) for img in p_element.iter(_VML_IMAGEDATA)]
    blobs = []
    for rid in rids:
        if not rid:
            continue
        part = doc.part.related_parts.get(rid)
        if part is not None and part.blob:
            blobs.append(part.blob)
    return blobs


def _split(doc: Document) -> Tuple[List[Tuple[str, str]], List[Tuple[str, bytes]]]:
    """Returns ([(section_name, markdown_body), ...], [(section_name, image_bytes), ...])
    with images in document order, tagged with whichever section was open
    when they were encountered."""
    sections: List[Tuple[str, List[str]]] = []
    images: List[Tuple[str, bytes]] = []

    for para in _iter_all_paragraphs(doc):
        for blob in _image_blobs_in(doc, para._p):
            if not sections:
                sections.append(("Preface", []))
            images.append((sections[-1][0], blob))

        text = para.text.strip()
        if not text or is_garbled(text) or is_page_number_artifact(text) or is_footer_url_artifact(text):
            # A "Page N" or "genai.owasp.org" footer typed inline as its own
            # paragraph rather than living in a real Word header/footer part
            # (observed in the real LLM Top 10 docx). Pagination and footer
            # branding from the English source's layout are meaningless —
            # and wrong — once reflowed into a translated document with a
            # different page count/template; Process 2's render_config.json
            # adds real page numbers and footer branding for final layout.
            continue
        style_name = para.style.name or ""
        if style_name == "Heading 1":
            sections.append((_slugify(text), [f"# {text}"]))
        elif style_name.lower().startswith(("heading", "subhead")):
            # Real-world exports use style names like "Heading 3 - No TOC" or
            # "Subhead" for sub-levels, not just a clean "Heading N" — treat
            # any of them as a subheading rather than losing the structure.
            last_token = style_name.split()[-1]
            level = "#" * min(int(last_token), 6) if last_token.isdigit() else "##"
            if not sections:
                sections.append(("Preface", []))
            sections[-1][1].append(f"{level} {text}")
        else:
            if not sections:
                sections.append(("Preface", []))
            sections[-1][1].append(text)

    # A section name can exist purely to anchor an image encountered before
    # any heading (e.g. "Preface") with no text ever appended to it — that's
    # not a translatable section, and an empty/whitespace-only message is
    # rejected outright by a real translation call (only ever surfaced in
    # production, never in --offline mode, since that skips the call).
    text_sections = [(name, "\n\n".join(lines) + "\n") for name, lines in sections if lines]
    return text_sections, images


def split_docx(docx_path: Path) -> List[Tuple[str, str]]:
    """Return [(section_name, markdown_body), ...] split on Heading 1 styles."""
    return _split(Document(str(docx_path)))[0]


def split_into_source(docx_path: Path, source_dir: Path) -> List[str]:
    """Write each split section as <source_dir>/<name>.md, plus a <name>.svg
    + images/<name>.png(+_original.png) per image with enough confidently-
    recognized text to be worth localizing. Returns all names written
    (sections and figures alike — translate_section.py treats both as
    sections to draft)."""
    doc = Document(str(docx_path))
    text_sections, image_candidates = _split(doc)

    written = []
    for name, body in text_sections:
        (source_dir / f"{name}.md").write_text(body)
        written.append(name)

    images_dir = source_dir / "images"
    figure_counts: dict = {}
    figures_by_section: dict = {}
    for section_name, blob in image_candidates:
        figure_counts[section_name] = figure_counts.get(section_name, 0) + 1
        suffix = "" if figure_counts[section_name] == 1 else f"_{figure_counts[section_name]}"
        name_hint = f"{section_name}_Figure{suffix}"
        href = f"images/{name_hint}.png"

        try:
            figure = image_svg.convert_image(blob, name_hint, href)
        except Exception:
            continue  # a malformed/unsupported embedded image is not fatal to the rest of the split
        if figure is None:
            continue

        images_dir.mkdir(exist_ok=True)
        (images_dir / f"{figure.name}.png").write_bytes(figure.base_png)
        (images_dir / f"{figure.name}_original.png").write_bytes(figure.original_png)
        (source_dir / f"{figure.name}.svg").write_text(figure.svg)
        written.append(figure.name)
        figures_by_section.setdefault(section_name, []).append(figure.name)

    raw_dir = source_dir / "_raw"
    raw_dir.mkdir(exist_ok=True)
    docx_path.rename(raw_dir / docx_path.name)

    # `written`'s order (all text, then all figures) suits translation, which
    # treats every name independently -- render.sh wants a figure where it
    # actually sat in the document, right after its enclosing section.
    manifest_order = []
    text_section_names = {name for name, _ in text_sections}
    # A section that anchored only an image and no text (a "Preface" opened
    # purely by a cover image before the first real heading exists) never
    # made it into text_sections -- filtered out as having no body -- but
    # its figure is still a real file. It can only precede every real
    # section (the auto-"Preface" only ever gets created before the first
    # heading is found), so flush any such orphans first, in encounter order.
    for orphan_section in list(figures_by_section):
        if orphan_section not in text_section_names:
            manifest_order.extend(figures_by_section.pop(orphan_section))
    for name, _ in text_sections:
        manifest_order.append(name)
        manifest_order.extend(figures_by_section.pop(name, []))
    write_manifest(source_dir, manifest_order)
    return written
